from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.drawing.image import Image as XLImage
from PIL import Image

from mdir import __version__
from mdir.workspace_app import (
    EditableWorkbookModel,
    _parse_clipboard_value,
    _pane_hotkey_action,
    _file_name_matches_filter,
    _search_arrow_target_index,
    _left_file_tab_action,
    _selection_workbook_action,
    _merge_recent_directories,
    _send_paths_to_recycle_bin,
)
from mdir.theme import effective_theme_name, normalize_theme_mode, palette_for
from mdir.excel_viewer import APP_VERSION as EXCEL_VIEWER_VERSION, AxisMetrics, _sheet_content_extent
from mdir.workspace_app import APP_VERSION as WORKSPACE_APP_VERSION
from mdir.links import (
    LinkDefinition,
    expand_link_text,
    load_links,
    load_mdir_links,
    parse_links,
    save_links,
)


class WorkspaceCoreTests(unittest.TestCase):
    def test_version_is_consistent_across_entry_points(self):
        self.assertEqual(WORKSPACE_APP_VERSION, __version__)
        self.assertEqual(EXCEL_VIEWER_VERSION, __version__)

    def test_theme_modes_and_palettes(self):
        self.assertEqual(normalize_theme_mode("light"), "Bright")
        self.assertEqual(normalize_theme_mode("DARK"), "Dark")
        self.assertEqual(normalize_theme_mode("anything-else"), "System")
        self.assertEqual(effective_theme_name("Bright"), "Bright")
        self.assertEqual(effective_theme_name("Dark"), "Dark")
        self.assertEqual(palette_for("Bright")["name"], "Bright")
        self.assertEqual(palette_for("Dark")["name"], "Dark")
        self.assertNotEqual(palette_for("Bright")["window"], palette_for("Dark")["window"])

    def test_recent_directory_history_is_mru_and_deduplicated(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            first = root / "first"
            second = root / "second"
            third = root / "third"
            for path in (first, second, third):
                path.mkdir()
            values = [first, second, first, third]
            recent = _merge_recent_directories(second, values, limit=3)
            self.assertEqual(recent, [second.resolve(), first.resolve(), third.resolve()])

    def test_filename_search_filter_is_case_insensitive_and_supports_multiple_terms(self):
        self.assertTrue(_file_name_matches_filter("010719-9 TJ.xlsx", "010719"))
        self.assertTrue(_file_name_matches_filter("010719-9 TJ.xlsx", "010719 tj"))
        self.assertTrue(_file_name_matches_filter("010719-9 TJ.xlsx", "TJ 9"))
        self.assertTrue(_file_name_matches_filter("ABC.XLSX", "abc"))
        self.assertTrue(_file_name_matches_filter("ABC.XLSX", ""))
        self.assertFalse(_file_name_matches_filter("010719-9 TJ.xlsx", "010620"))
        self.assertFalse(_file_name_matches_filter("010719-9 TJ.xlsx", "010719 win"))

    def test_left_file_list_tab_is_blocked(self):
        self.assertEqual(_left_file_tab_action(), "break")

    def test_search_arrow_leaves_search_and_targets_filtered_rows(self):
        self.assertEqual(_search_arrow_target_index(5, None, 1), 0)
        self.assertEqual(_search_arrow_target_index(5, None, -1), 4)
        self.assertEqual(_search_arrow_target_index(5, 1, 1), 2)
        self.assertEqual(_search_arrow_target_index(5, 3, -1), 2)
        self.assertEqual(_search_arrow_target_index(5, 0, -1), 0)
        self.assertEqual(_search_arrow_target_index(5, 4, 1), 4)
        self.assertIsNone(_search_arrow_target_index(0, None, 1))

    def test_single_left_selection_decides_right_workbook_action(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            folder = root / "folder"
            folder.mkdir()
            excel = root / "book.xlsx"
            excel.write_bytes(b"placeholder")
            text = root / "notes.txt"
            text.write_text("hello", encoding="utf-8")

            self.assertEqual(_selection_workbook_action([excel]), "load")
            self.assertEqual(_selection_workbook_action([folder]), "clear")
            self.assertEqual(_selection_workbook_action([text]), "clear")
            self.assertEqual(_selection_workbook_action([]), "keep")
            self.assertEqual(_selection_workbook_action([excel, text]), "keep")

    def test_parse_clipboard_value(self):
        self.assertEqual(_parse_clipboard_value("12"), 12)
        self.assertEqual(_parse_clipboard_value("12.5"), 12.5)
        self.assertEqual(_parse_clipboard_value("=A1+B1"), "=A1+B1")
        self.assertEqual(_parse_clipboard_value("00123"), "00123")

    def test_link_parser_keeps_mdir_types_and_pane(self):
        values = [
            {"label": "CHOIS", "type": "folder", "target": r"D:\\sys_back\\chois", "pane": "left"},
            {"label": "GitHub", "type": "web", "target": "https://example.com"},
            {"label": "CMD", "type": "action", "target": "powershell_here", "pane": "active"},
            {"label": "Tool", "type": "program", "target": "tool.exe", "args": ["-x"]},
        ]
        links = parse_links(values)
        self.assertEqual([item.kind for item in links], ["folder", "web", "action", "program"])
        self.assertEqual(links[0].pane, "left")
        self.assertEqual(links[3].args, ("-x",))

    def test_links_roundtrip_mdir_compatible_schema(self):
        with TemporaryDirectory() as td:
            path = Path(td) / "links.json"
            links = [
                LinkDefinition("Home", "folder", "{home}"),
                LinkDefinition("GitHub", "web", "https://example.com"),
                LinkDefinition("CMD", "action", "powershell_here"),
                LinkDefinition("Tool", "program", "tool.exe", ("-x",), "right"),
            ]
            save_links(links, path)
            self.assertEqual(load_links(path, Path(td) / "missing-mdir.json"), links)

    def test_load_mdir_keeps_non_folder_items(self):
        with TemporaryDirectory() as td:
            path = Path(td) / "mdir-links.json"
            path.write_text(
                '[{"label":"Work","type":"folder","target":"C:/Work"},'
                '{"label":"GitHub","type":"web","target":"https://example.com"},'
                '{"label":"CMD","type":"action","target":"powershell_here"}]',
                encoding="utf-8",
            )
            links = load_mdir_links(path)
            self.assertEqual([item.kind for item in links], ["folder", "web", "action"])

    def test_legacy_folder_config_migrates_to_full_mdir_list(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            xexcel = root / "xexcel.json"
            mdir = root / "mdir.json"
            xexcel.write_text(
                '[{"label":"Home","target":"{home}"},'
                '{"label":"SPEC","target":"C:/SPEC"}]',
                encoding="utf-8",
            )
            mdir.write_text(
                '[{"label":"Home","type":"folder","target":"{home}"},'
                '{"label":"SPEC","type":"folder","target":"C:/SPEC"},'
                '{"label":"GitHub","type":"web","target":"https://example.com"},'
                '{"label":"CMD","type":"action","target":"powershell_here"}]',
                encoding="utf-8",
            )
            links = load_links(xexcel, mdir)
            self.assertEqual([x.label for x in links], ["Home", "SPEC", "GitHub", "CMD"])

    def test_expand_link_text_supports_mdir_tokens(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            current = root / "current"
            selected = current / "book.xlsx"
            workbook = root / "wb" / "open.xlsx"
            expanded = expand_link_text(
                "{current}|{selected}|{right_selected}|{home}|{project}",
                current=current, selected=selected, workbook=workbook, project=root / "project",
            )
            self.assertIn(str(current), expanded)
            self.assertIn(str(selected), expanded)
            self.assertIn(str(workbook), expanded)
            self.assertIn(str(root / "project"), expanded)

    def test_pane_hotkey_action_windows_and_tk_masks(self):
        self.assertEqual(_pane_hotkey_action("Left", 37, 0x0004), "left")
        self.assertEqual(_pane_hotkey_action("Right", 39, 0x0004), "right")
        self.assertEqual(_pane_hotkey_action("1", 49, 0x20000), "left")
        self.assertEqual(_pane_hotkey_action("2", 50, 0x20000), "right")
        self.assertEqual(_pane_hotkey_action("Left", 37, 0, ctrl_down=True), "left")
        self.assertEqual(_pane_hotkey_action("2", 50, 0, alt_down=True), "right")
        self.assertIsNone(_pane_hotkey_action("Left", 37, 0))
        self.assertIsNone(_pane_hotkey_action("1", 49, 0x20000, alt_down=False))
        self.assertIsNone(_pane_hotkey_action("2", 50, 0x20000, alt_down=False))
        self.assertIsNone(_pane_hotkey_action("3", 51, 0x20000, alt_down=False))
        self.assertEqual(_pane_hotkey_action("1", 49, 0x20000), "left")

    def test_edit_and_save_copy(self):
        with TemporaryDirectory() as td:
            source = Path(td) / "source.xlsx"
            target = Path(td) / "target.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "Original"
            wb.save(source)

            model = EditableWorkbookModel(source)
            self.assertTrue(model.editable)
            sheet = model.sheets[0]
            sheet.set_value(1, 1, "Changed")
            sheet.set_value(2, 2, 123)
            self.assertTrue(model.dirty)
            model.save(target, create_backup=False)
            model.close()

            check = load_workbook(target, data_only=False)
            self.assertEqual(check.active["A1"].value, "Changed")
            self.assertEqual(check.active["B2"].value, 123)
            check.close()

    def test_embedded_image_survives_edit_and_save(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            image_path = root / "product.png"
            Image.new("RGB", (40, 30), (120, 80, 40)).save(image_path)
            source = root / "image-source.xlsx"
            target = root / "image-target.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "Product"
            ws.add_image(XLImage(str(image_path)), "B2")
            wb.save(source)

            model = EditableWorkbookModel(source)
            self.assertEqual(len(model.sheets[0].images), 1)
            model.sheets[0].set_value(1, 1, "Changed Product")
            model.save(target, create_backup=False)
            model.close()

            check = load_workbook(target)
            self.assertEqual(check.active["A1"].value, "Changed Product")
            self.assertEqual(len(check.active._images), 1)
            check.close()

    def test_backup_on_original_overwrite(self):
        with TemporaryDirectory() as td:
            source = Path(td) / "source.xlsx"
            wb = Workbook()
            wb.active["A1"] = "Before"
            wb.save(source)
            model = EditableWorkbookModel(source)
            model.sheets[0].set_value(1, 1, "After")
            model.save()
            backups = list((Path(td) / "xExcel_Backup").glob("source-*.xlsx"))
            self.assertEqual(len(backups), 1)
            model.close()

    def test_excel_layout_fidelity_metadata(self):
        with TemporaryDirectory() as td:
            source = Path(td) / "layout.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.sheet_view.showGridLines = False
            ws.sheet_format.defaultRowHeight = 15
            ws.column_dimensions["B"].hidden = True
            ws.column_dimensions["C"].hidden = True
            ws.column_dimensions["D"].width = 20
            ws.row_dimensions[2].hidden = True
            ws.row_dimensions[4].height = 30
            ws.merge_cells("A1:D1")
            ws["A1"] = "ITEM#"
            ws["A1"].font = Font(name="Calibri", size=16, bold=True)
            ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
            ws["A1"].fill = PatternFill("solid", fgColor="FFF2CC")
            thin = Side(style="thin", color="000000")
            ws["A5"] = "Bordered"
            ws["A5"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
            wb.save(source)

            model = EditableWorkbookModel(source)
            sheet = model.sheets[0]
            cols = sheet.column_overrides(1.0)
            rows = sheet.row_overrides(1.0)
            self.assertFalse(sheet.show_gridlines)
            self.assertEqual(cols[2], 0)
            self.assertEqual(cols[3], 0)
            self.assertGreater(cols[4], 100)
            self.assertEqual(rows[2], 0)
            self.assertGreater(rows[4], 30)
            style = sheet.style(1, 1)
            self.assertTrue(style.bold)
            self.assertEqual(style.font_name, "Calibri")
            self.assertEqual(style.vertical, "center")
            self.assertEqual(style.fill.upper(), "#FFF2CC")
            bordered = sheet.style(5, 1)
            self.assertEqual(bordered.left.style, "thin")
            self.assertEqual(sheet.merged_ranges[0], (1, 1, 4, 1))
            model.close()

    def test_scroll_extent_includes_image_below_used_cells(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            image_path = root / "tall.png"
            Image.new("RGB", (120, 700), (180, 180, 180)).save(image_path)
            source = root / "tall-image.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "Only used cell"
            picture = XLImage(str(image_path))
            picture.width = 120
            picture.height = 700
            ws.add_image(picture, "B2")
            wb.save(source)

            model = EditableWorkbookModel(source)
            sheet = model.sheets[0]
            rows = AxisMetrics(sheet.max_row, sheet.default_row_size(1.0), sheet.row_overrides(1.0))
            cols = AxisMetrics(sheet.max_column, sheet.default_column_size(1.0), sheet.column_overrides(1.0))
            width, height = _sheet_content_extent(sheet, rows, cols, 1.0)
            self.assertGreater(height, rows.total + 500)
            self.assertGreaterEqual(width, cols.total)
            model.close()

    def test_all_worksheet_names_are_loaded(self):
        with TemporaryDirectory() as td:
            source = Path(td) / "many-sheets.xlsx"
            wb = Workbook()
            wb.active.title = "Overview"
            for name in ("Costing", "Images", "Buyer Notes", "Order", "Archive"):
                wb.create_sheet(name)
            wb.save(source)
            model = EditableWorkbookModel(source)
            self.assertEqual(
                [sheet.title for sheet in model.sheets],
                ["Overview", "Costing", "Images", "Buyer Notes", "Order", "Archive"],
            )
            model.close()

    def test_enter_binding_opens_default_application(self):
        import inspect
        from mdir.workspace_app import ExcelWorkspaceApp
        source = inspect.getsource(ExcelWorkspaceApp._file_enter_open_default)
        self.assertIn("_open_path_with_default_application", source)
        self.assertIn("path.is_dir()", source)

    def test_multiselect_and_delete_binding_are_enabled(self):
        import inspect
        from mdir.workspace_app import ExcelWorkspaceApp
        source = inspect.getsource(ExcelWorkspaceApp._build_ui)
        self.assertIn('selectmode="extended"', source)
        self.assertIn('bind("<Delete>", self._delete_selected_items)', source)

    def test_recycle_helper_never_permanently_deletes(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            one = root / "one.xlsx"
            two = root / "two.xlsx"
            one.write_text("1", encoding="utf-8")
            two.write_text("2", encoding="utf-8")
            captured = []
            with patch("mdir.workspace_app.send2trash", side_effect=lambda p: captured.append(Path(p))):
                moved, failed = _send_paths_to_recycle_bin([one, two])
            self.assertEqual(moved, [one, two])
            self.assertEqual(failed, [])
            self.assertEqual(captured, [one, two])
            self.assertTrue(one.exists())
            self.assertTrue(two.exists())

    def test_right_mouse_drag_selection_bindings_are_enabled(self):
        import inspect
        from mdir.workspace_app import ExcelWorkspaceApp
        build = inspect.getsource(ExcelWorkspaceApp._build_ui)
        self.assertIn('bind("<ButtonPress-3>", self._file_right_drag_start)', build)
        self.assertIn('bind("<B3-Motion>", self._file_right_drag_motion)', build)
        self.assertIn('bind("<ButtonRelease-3>", self._file_right_drag_end)', build)
        process = inspect.getsource(ExcelWorkspaceApp._file_right_drag_process_y)
        self.assertIn('_right_drag_seen', inspect.getsource(ExcelWorkspaceApp._file_right_drag_toggle_iid))
        self.assertIn('range(previous + step, index + step, step)', process)

    def test_right_drag_keeps_starting_selected_row_and_adds_crossed_rows(self):
        from mdir.workspace_app import ExcelWorkspaceApp

        class FakeTree:
            def __init__(self):
                self.selected = {"row2"}
                self.focused = None
            def selection(self):
                return tuple(self.selected)
            def selection_add(self, iid):
                self.selected.add(iid)
            def selection_remove(self, iid):
                self.selected.discard(iid)
            def focus(self, iid):
                self.focused = iid
            def see(self, iid):
                pass

        app = ExcelWorkspaceApp.__new__(ExcelWorkspaceApp)
        app.file_tree = FakeTree()
        app._right_drag_seen = set()

        app._file_right_drag_toggle_iid("row2")
        app._file_right_drag_toggle_iid("row3")
        app._file_right_drag_toggle_iid("row4")
        self.assertEqual(app.file_tree.selected, {"row2", "row3", "row4"})
        self.assertEqual(app.file_tree.focused, "row4")


if __name__ == "__main__":
    unittest.main()
